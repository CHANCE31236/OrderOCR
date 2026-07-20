from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from src.business.validators import duplicate_product_codes, safe_order_filename, validate_product_code
from src.ocr.schemas import ReviewedRow


class ExportBlockedError(ValueError):
    pass


def export_order_excel(
    order_number: str,
    rows: list[ReviewedRow],
    output_directory: str | Path,
    *,
    allow_duplicates: bool = False,
) -> Path:
    safe_order_filename(order_number)
    if not rows:
        raise ExportBlockedError("There are no product rows to export")
    pending = [row.observation.row_number for row in rows if row.needs_review and not row.manually_confirmed]
    if pending:
        raise ExportBlockedError(f"These rows still require manual confirmation: {pending}")
    codes = [validate_product_code(row.observation.product_code) for row in rows]
    duplicates = duplicate_product_codes(codes)
    if duplicates and not allow_duplicates:
        raise ExportBlockedError("Duplicate product codes require confirmation: " + ", ".join(duplicates))
    for row in rows:
        if row.final_box_count is None:
            raise ExportBlockedError(f"The final box count for row {row.observation.row_number} is unknown")
        if not isinstance(row.final_box_count, int) or isinstance(row.final_box_count, bool) or row.final_box_count < 0:
            raise ExportBlockedError("The box count must be a non-negative integer")

    directory = Path(output_directory)
    try:
        directory.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ExportBlockedError(f"Could not create the output folder: {exc}") from exc
    target = directory / f"{order_number}.xlsx"
    temporary = directory / f".{order_number}.tmp.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Order"
    sheet.append(["Product code", "Box count"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.observation.product_code, row.final_box_count])
        sheet.cell(sheet.max_row, 1).number_format = "@"
        sheet.cell(sheet.max_row, 2).number_format = "0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.column_dimensions["A"].width = min(60, max(14, max(len(code) for code in codes) + 3))
    sheet.column_dimensions["B"].width = 12
    try:
        workbook.save(temporary)
        temporary.replace(target)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise ExportBlockedError("The Excel file may be open, or the output folder may not be writable") from exc
    return target
