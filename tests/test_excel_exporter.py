from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.export.excel_exporter import ExportBlockedError, export_order_excel
from src.ocr.schemas import ReviewedRow, RowObservation


def make_row(code: str, box: int, *, needs_review: bool = False, confirmed: bool = False) -> ReviewedRow:
    observation = RowObservation(
        row_number=1,
        product_code=code,
        printed_box_count=box,
        printed_total_quantity=150,
        circle_status="full_circle",
        handwritten_override=None,
        override_is_clearly_associated=False,
        calculated_final_box_count=box,
        product_code_confidence=0.99,
        circle_confidence=0.99,
        overall_confidence=0.99,
        needs_review=needs_review,
        reason="Closed circle",
    )
    return ReviewedRow(observation=observation, final_box_count=box, needs_review=needs_review, manually_confirmed=confirmed, status="auto_confirmed")


def test_excel_has_only_two_columns_and_preserves_leading_zero(tmp_path: Path):
    path = export_order_excel("Vs20260713-12", [make_row("00001535", 5)], tmp_path)
    book = load_workbook(path)
    sheet = book.active
    assert sheet.max_column == 2
    assert [sheet["A1"].value, sheet["B1"].value] == ["Product code", "Box count"]
    assert sheet["A2"].value == "00001535"
    assert isinstance(sheet["A2"].value, str)
    assert sheet["A2"].number_format == "@"
    assert sheet["B2"].value == 5
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B2"


def test_unreviewed_uncertain_row_is_blocked(tmp_path: Path):
    with pytest.raises(ExportBlockedError, match="manual confirmation"):
        export_order_excel("A-1", [make_row("X", 1, needs_review=True)], tmp_path)


def test_duplicate_codes_require_explicit_confirmation(tmp_path: Path):
    rows = [make_row("DUP", 1), make_row("DUP", 2)]
    with pytest.raises(ExportBlockedError, match="Duplicate product codes"):
        export_order_excel("A-1", rows, tmp_path)
    assert export_order_excel("A-1", rows, tmp_path, allow_duplicates=True).exists()


def test_special_product_code_characters_are_not_changed(tmp_path: Path):
    codes = ["DV1731-01", "ORCHIDEE RIZ LONG 5x4.5kg", "FT7850-1/107154"]
    path = export_order_excel("A-2", [make_row(code, 1) for code in codes], tmp_path)
    sheet = load_workbook(path).active
    assert [sheet.cell(row, 1).value for row in range(2, 5)] == codes


@pytest.mark.parametrize("code", ['=HYPERLINK("https://example.com", "code")', "=1+1", "#N/A", "+ABC", "@ABC"])
def test_product_codes_are_stored_as_literal_text(tmp_path: Path, code: str):
    path = export_order_excel("A-3", [make_row(code, 1)], tmp_path)
    book = load_workbook(path)
    try:
        assert book.active["A2"].value == code
        assert book.active["A2"].data_type == "s"
    finally:
        book.close()
